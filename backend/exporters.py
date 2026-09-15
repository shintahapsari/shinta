"""Server-side PDF & Excel export."""
import io
from datetime import datetime, timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

UNEJ_BLUE = colors.HexColor("#0B2545")
UNEJ_GOLD = colors.HexColor("#F5A623")

INSTITUTION = ("Sistem Terintegrasi Kemitraan Teknologi Industri Pertanian — "
               "Program Studi Teknologi Industri Pertanian, Fakultas Teknologi Pertanian, "
               "Universitas Jember")


def to_excel(title: str, headers: list, rows: list) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Laporan"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(headers)))
    c = ws.cell(row=1, column=1, value=title)
    c.font = Font(bold=True, size=13, color="0B2545")
    c.alignment = Alignment(horizontal="left")
    ws.cell(row=2, column=1, value=INSTITUTION).font = Font(size=9, italic=True, color="475569")
    hdr_fill = PatternFill("solid", fgColor="0B2545")
    for i, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=i, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center")
    for r, row in enumerate(rows, start=5):
        for i, val in enumerate(row, start=1):
            ws.cell(row=r, column=i, value=val)
    for i, h in enumerate(headers, start=1):
        maxlen = max([len(str(h))] + [len(str(row[i - 1])) for row in rows]) if rows else len(str(h))
        ws.column_dimensions[ws.cell(row=4, column=i).column_letter].width = min(max(12, maxlen + 3), 50)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(title: str, headers: list, rows: list) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), topMargin=1.2 * cm,
                            bottomMargin=1.2 * cm, leftMargin=1.2 * cm, rightMargin=1.2 * cm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("t", parent=styles["Title"], fontSize=15, textColor=UNEJ_BLUE, alignment=0)
    sub_style = ParagraphStyle("s", parent=styles["Normal"], fontSize=8, textColor=colors.HexColor("#475569"))
    cell_style = ParagraphStyle("c", parent=styles["Normal"], fontSize=8, leading=10)
    hdr_style = ParagraphStyle("h", parent=styles["Normal"], fontSize=8, leading=10, textColor=colors.white, fontName="Helvetica-Bold")
    elems = [Paragraph(title, title_style), Spacer(1, 4), Paragraph(INSTITUTION, sub_style),
             Paragraph(f"Digenerate: {datetime.now(timezone.utc).strftime('%d %b %Y %H:%M UTC')}", sub_style),
             Spacer(1, 12)]
    data = [[Paragraph(str(h), hdr_style) for h in headers]]
    for row in rows:
        data.append([Paragraph(str(v), cell_style) for v in row])
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), UNEJ_BLUE),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LINEBELOW", (0, 0), (-1, 0), 1.5, UNEJ_GOLD),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    elems.append(table)
    doc.build(elems)
    return buf.getvalue()
